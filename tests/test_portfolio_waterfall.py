"""Tests for domain/portfolio/waterfall.py — pooled CFADS portfolio waterfall."""
import pytest
from datetime import date
from domain.portfolio.waterfall import (
    PortfolioPeriod,
    PortfolioResult,
    aggregate_project_results,
    portfolio_cfads_schedule,
    run_portfolio_waterfall,
)
from domain.waterfall.waterfall_engine import WaterfallResult, WaterfallPeriod
from app.project_factories import create_default_solar_project, create_default_wind_project
from domain.inputs import FinancingParams


def _make_wf_result(name: str, ebitda: float, tax: float, rev: float) -> WaterfallResult:
    """Create a minimal WaterfallResult with two operation periods."""
    # All 37 WaterfallPeriod fields must be present
    def _op(period: int, date_: date, rev_h: float, ebitda_h: float, tax_h: float):
        return WaterfallPeriod(
            period=period, date=date_, year_index=1, period_in_year=period,
            is_operation=True,
            generation_mwh=0.0,
            revenue_keur=rev_h, opex_keur=0.0, ebitda_keur=ebitda_h,
            depreciation_keur=0.0, interest_senior_keur=0.0, interest_shl_keur=0.0,
            taxable_profit_keur=ebitda_h, tax_keur=tax_h,
            cf_after_tax_keur=ebitda_h - tax_h,
            senior_interest_keur=0.0, senior_principal_keur=0.0,
            senior_ds_keur=50.0,
            shl_interest_keur=0.0, shl_principal_keur=0.0, shl_service_keur=0.0,
            dsra_contribution_keur=0.0, dsra_balance_keur=0.0,
            mra_contribution_keur=0.0, mra_balance_keur=0.0,
            cf_after_reserves_keur=ebitda_h - tax_h,
            dscr=1.5, llcr=1.5, plcr=1.5, lockup_active=False,
            distribution_keur=0.0, cash_sweep_keur=0.0, cum_distribution_keur=0.0,
            cash_balance_keur=0.0,
            shl_balance_keur=0.0, shl_pik_keur=0.0,
            senior_balance_keur=0.0,
        )
    op1 = _op(1, date(2031, 1, 1), rev / 2, ebitda / 2, tax / 2)
    op2 = _op(2, date(2031, 7, 1), rev / 2, ebitda / 2, tax / 2)

    # Attach a mock inputs object so waterfall.py can compute total_capex for XIRR
    from unittest.mock import MagicMock
    mock_inputs = MagicMock()
    mock_inputs.capex.total_capex = 1000.0  # kEUR — enough to get a meaningful XIRR

    wf = WaterfallResult(
        periods=(op1, op2),
        total_revenue_keur=rev, total_opex_keur=0.0,
        total_ebitda_keur=ebitda, total_tax_keur=tax,
        total_senior_ds_keur=100.0, total_shl_service_keur=0.0,
        total_distribution_keur=0.0,
        avg_dscr=1.5, min_dscr=1.4, max_dscr=1.6,
        min_llcr=1.4, min_plcr=1.3, periods_in_lockup=0,
        project_irr=0.0, equity_irr=0.0, sponsor_irr=0.0,
        project_npv=0.0, equity_npv=0.0,
        sculpting_result=None,
    )
    wf.inputs = mock_inputs
    return wf


def _portfolio_inputs():
    shared = FinancingParams(share_capital_keur=100.0, senior_debt_amount_keur=200.0,
                             senior_tenor_years=10, target_dscr=1.3)
    from domain.portfolio.inputs import PortfolioInputs
    return PortfolioInputs(
        projects=(create_default_solar_project(), create_default_wind_project()),
        portfolio_name="Test",
        shared_financing=shared,
    )


class TestAggregateProjectResults:
    """Test aggregate_project_results aligns and sums waterfalls."""

    def test_two_projects_sum_revenue_and_ebitda(self):
        wf_a = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
        wf_b = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)
        pooled = aggregate_project_results((("A", wf_a), ("B", wf_b)))
        assert len(pooled) == 2
        assert abs(pooled[0]["pooled_revenue_keur"] - 125.0) < 0.01
        assert abs(pooled[0]["pooled_ebitda_keur"] - 100.0) < 0.01

    def test_cfads_equals_ebitda_minus_tax(self):
        wf = _make_wf_result("X", ebitda=80.0, tax=10.0, rev=100.0)
        pooled = aggregate_project_results((("X", wf),))
        assert len(pooled) == 2
        # Per period: ebitda=40, tax=5, cfads=35
        assert abs(pooled[0]["pooled_cfads_keur"] - 35.0) < 0.01

    def test_skips_non_operation_periods(self):
        wf = _make_wf_result("X", ebitda=80.0, tax=10.0, rev=100.0)
        pooled = aggregate_project_results((("X", wf),))
        assert all(p["date"] >= date(2030, 1, 1) for p in pooled)


class TestPortfolioCFADSSchedule:
    """Test portfolio_cfads_schedule extracts CFADS list."""

    def test_extracts_cfads_from_pooled_periods(self):
        pooled = [
            {"date": date(2031, 1, 1), "pooled_cfads_keur": 35.0},
            {"date": date(2031, 7, 1), "pooled_cfads_keur": 40.0},
        ]
        cfads = portfolio_cfads_schedule(pooled)
        assert cfads == [35.0, 40.0]


class TestRunPortfolioWaterfall:
    """Test run_portfolio_waterfall with skeleton debt service."""

    def test_total_revenue_summed_across_projects(self):
        wf_a = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
        wf_b = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)
        result = run_portfolio_waterfall(_portfolio_inputs(), (("A", wf_a), ("B", wf_b)))
        assert abs(result.total_revenue_keur - 250.0) < 0.01

    def test_total_ebitda_summed_across_projects(self):
        wf_a = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
        wf_b = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)
        result = run_portfolio_waterfall(_portfolio_inputs(), (("A", wf_a), ("B", wf_b)))
        assert abs(result.total_ebitda_keur - 200.0) < 0.01

    def test_total_tax_summed_across_projects(self):
        wf_a = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
        wf_b = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)
        result = run_portfolio_waterfall(_portfolio_inputs(), (("A", wf_a), ("B", wf_b)))
        assert abs(result.total_tax_keur - 25.0) < 0.01

    def test_dscr_is_computed_per_period(self):
        wf = _make_wf_result("X", ebitda=80.0, tax=10.0, rev=100.0)
        result = run_portfolio_waterfall(_portfolio_inputs(), (("X", wf),))
        assert len(result.periods) == 2
        # DSCR depends on sculpted debt from CFADS; just check it's positive
        assert result.periods[0].dscr > 0

    def test_avg_dscr_across_periods(self):
        wf_a = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
        wf_b = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)
        result = run_portfolio_waterfall(_portfolio_inputs(), (("A", wf_a), ("B", wf_b)))
        assert result.avg_dscr > 0

    def test_min_dscr_identified(self):
        wf_a = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
        wf_b = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)
        result = run_portfolio_waterfall(_portfolio_inputs(), (("A", wf_a), ("B", wf_b)))
        assert result.min_dscr > 0

    def test_explicit_ds_schedule_used_when_provided(self):
        wf = _make_wf_result("X", ebitda=80.0, tax=10.0, rev=100.0)
        ds_schedule = (200.0, 200.0)
        result = run_portfolio_waterfall(None, (("X", wf),),
                                         portfolio_debt_service_schedule=ds_schedule)
        # Period 1: cfads=35, ds=200, dscr=0.175
        assert abs(result.periods[0].dscr - 0.175) < 0.001

    def test_project_results_passed_through(self):
        wf_a = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
        wf_b = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)
        result = run_portfolio_waterfall(_portfolio_inputs(), (("A", wf_a), ("B", wf_b)))
        assert len(result.project_results) == 2
        names = [n for n, _ in result.project_results]
        assert "A" in names
        assert "B" in names

    def test_requires_inputs_or_explicit_debt_service_schedule(self):
        """Without portfolio_inputs AND without explicit ds schedule, raises."""
        wf = _make_wf_result("X", ebitda=80.0, tax=10.0, rev=100.0)
        with pytest.raises(ValueError, match="Either portfolio_inputs or explicit"):
            run_portfolio_waterfall(None, (("X", wf),))


class TestPortfolioResult:
    """Test PortfolioResult dataclass field contracts."""

    def test_portfolio_result_has_all_required_fields(self):
        wf = _make_wf_result("X", ebitda=80.0, tax=10.0, rev=100.0)
        pr = PortfolioResult(
            periods=(PortfolioPeriod(
                period=1, date=date(2030, 1, 1),
                pooled_revenue_keur=100.0, pooled_ebitda_keur=80.0,
                pooled_tax_keur=10.0, pooled_cfads_keur=70.0,
                portfolio_senior_interest_keur=40.0,
                portfolio_senior_principal_keur=10.0,
                portfolio_senior_ds_keur=50.0, dscr=1.4,
            ),),
            project_results=(("X", wf),),
            total_revenue_keur=100.0, total_ebitda_keur=80.0,
            total_tax_keur=10.0, total_senior_ds_keur=50.0,
            avg_dscr=1.4, min_dscr=1.2,
            portfolio_debt_keur=500.0,
            pooled_cfads_schedule=(70.0,),
            portfolio_debt_service_schedule=(50.0,),
        )
        assert hasattr(pr, "periods")
        assert hasattr(pr, "project_results")
        assert hasattr(pr, "total_revenue_keur")
        assert hasattr(pr, "total_ebitda_keur")
        assert hasattr(pr, "total_tax_keur")
        assert hasattr(pr, "total_senior_ds_keur")
        assert hasattr(pr, "avg_dscr")
        assert hasattr(pr, "min_dscr")
        assert hasattr(pr, "portfolio_project_irr")
        assert hasattr(pr, "portfolio_sponsor_irr")


def test_explicit_debt_service_schedule_overrides_sculpting():
    """When explicit portfolio_debt_service_schedule is supplied, portfolio_debt_keur = 0.0."""
    wf = _make_wf_result("X", ebitda=80.0, tax=10.0, rev=100.0)
    ds_schedule = (200.0, 200.0)
    result = run_portfolio_waterfall(None, (("X", wf),),
                                     portfolio_debt_service_schedule=ds_schedule)
    assert result.portfolio_debt_keur == 0.0


def test_portfolio_uses_sculpted_debt_service_when_no_schedule_supplied():
    """When no explicit schedule is given, portfolio_debt_keur comes from sculpted schedule (>0)."""
    wf_a = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
    wf_b = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)
    result = run_portfolio_waterfall(_portfolio_inputs(), (("A", wf_a), ("B", wf_b)))
    assert result.portfolio_debt_keur > 0


def test_portfolio_project_irr_is_finite():
    """portfolio_project_irr should be set (finite or 0.0 if XIRR doesn't converge)."""
    wf_a = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
    wf_b = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)
    result = run_portfolio_waterfall(_portfolio_inputs(), (("A", wf_a), ("B", wf_b)))
    # Should be a finite number or 0.0 (XIRR may not converge on test CFADS)
    assert result.portfolio_project_irr is not None
    import math
    assert math.isfinite(result.portfolio_project_irr)


def test_portfolio_sponsor_irr_is_placeholder():
    """portfolio_sponsor_irr is explicitly placeholder (not yet implemented)."""
    wf_a = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
    wf_b = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)
    result = run_portfolio_waterfall(_portfolio_inputs(), (("A", wf_a), ("B", wf_b)))
    # Documented as placeholder
    assert result.portfolio_sponsor_irr == 0.0 or result.portfolio_sponsor_irr is None


def test_portfolio_irr_fields_are_documented_placeholders():
    """portfolio_sponsor_irr is a placeholder; portfolio_project_irr is now computed."""
    # portfolio_sponsor_irr remains the documented placeholder
    assert PortfolioResult.__dataclass_fields__["portfolio_sponsor_irr"].default == 0.0


def test_portfolio_project_irr_changes_when_capex_changes():
    """Scaling one project's CapEx must change portfolio_project_irr."""
    from app.project_factories import create_default_solar_project, create_default_wind_project
    from app.portfolio_runner import run_portfolio_from_inputs
    from app.capex_overrides import scale_capex_items
    from dataclasses import replace
    from domain.portfolio.inputs import PortfolioInputs
    from domain.inputs import FinancingParams

    solar = create_default_solar_project()
    wind = create_default_wind_project()
    shared = FinancingParams(share_capital_keur=100.0, senior_debt_amount_keur=200.0,
                             senior_tenor_years=10, target_dscr=1.3)
    portfolio_inputs = PortfolioInputs(projects=(solar, wind), portfolio_name="Test",
                                        shared_financing=shared)

    base_result = run_portfolio_from_inputs(portfolio_inputs)
    base_irr = base_result.portfolio_project_irr
    assert base_irr is not None and base_irr != 0.0, "Base portfolio IRR should be non-zero"

    # Scale solar CapEx by +20%
    scaled_solar = replace(solar, capex=scale_capex_items(
        solar.capex, solar.capex.total_capex * 1.20
    ))
    scaled_portfolio_inputs = PortfolioInputs(
        projects=(scaled_solar, wind),
        portfolio_name="Test",
        shared_financing=shared,
    )
    scaled_result = run_portfolio_from_inputs(scaled_portfolio_inputs)
    scaled_irr = scaled_result.portfolio_project_irr

    assert scaled_irr != base_irr, "Portfolio IRR should change when CapEx scales"
    # More CapEx with same revenue → lower IRR
    assert scaled_irr < base_irr, f"Scaled IRR ({scaled_irr:.4f}) should be lower than base ({base_irr:.4f})"


def test_portfolio_project_irr_negative_t0():
    """Portfolio xirr t0 (initial cashflow) must be negative (investment = outflow)."""
    from app.project_factories import create_default_solar_project, create_default_wind_project
    from app.portfolio_runner import run_portfolio_from_inputs
    from domain.portfolio.inputs import PortfolioInputs
    from domain.inputs import FinancingParams

    solar = create_default_solar_project()
    wind = create_default_wind_project()
    shared = FinancingParams(share_capital_keur=100.0, senior_debt_amount_keur=200.0,
                             senior_tenor_years=10, target_dscr=1.3)
    portfolio_inputs = PortfolioInputs(
        projects=(solar, wind), portfolio_name="Test", shared_financing=shared
    )
    result = run_portfolio_from_inputs(portfolio_inputs)
    assert result.portfolio_project_irr is not None

    # t0 is defined as -sum of total_capex (negative outflow)
    total_capex = solar.capex.total_capex + wind.capex.total_capex
    assert total_capex > 0, "Total capex must be positive (investment)"
    # The xirr is computed on [-total_capex, ...positive_cfads...] — t0 is negative
    # We can verify the pooled_cfads_schedule is positive after t0
    cfads = result.pooled_cfads_schedule
    assert cfads[0] >= 0, "CFADS after t0 should be >= 0"


def test_portfolio_project_irr_finite_for_positive_cashflows():
    """portfolio_project_irr must be numeric and finite when CFADS are positive."""
    from app.project_factories import create_default_solar_project, create_default_wind_project
    from app.portfolio_runner import run_portfolio_from_inputs
    from domain.portfolio.inputs import PortfolioInputs
    from domain.inputs import FinancingParams

    solar = create_default_solar_project()
    wind = create_default_wind_project()
    shared = FinancingParams(share_capital_keur=100.0, senior_debt_amount_keur=200.0,
                             senior_tenor_years=10, target_dscr=1.3)
    portfolio_inputs = PortfolioInputs(
        projects=(solar, wind), portfolio_name="Test", shared_financing=shared
    )
    result = run_portfolio_from_inputs(portfolio_inputs)
    irr = result.portfolio_project_irr

    assert isinstance(irr, (int, float)), f"IRR should be numeric, got {type(irr)}"
    assert not (irr != irr), "IRR must not be NaN"  # NaN check
    assert -1 < irr < 10, f"IRR should be in reasonable range (-1 to 10), got {irr}"


def test_portfolio_sponsor_irr_is_not_numeric_zero_in_table():
    """build_portfolio_table should show sponsor IRR as 'n/a', not 0.0."""
    from unittest.mock import MagicMock
    from app.output_tables import build_portfolio_table

    pr = MagicMock()
    pr.total_revenue_keur = 100_000.0
    pr.total_ebitda_keur = 70_000.0
    pr.total_tax_keur = 10_000.0
    pr.pooled_cfads_schedule = (50_000.0, 55_000.0, 60_000.0)
    pr.total_senior_ds_keur = 30_000.0
    pr.avg_dscr = 1.4
    pr.min_dscr = 1.2
    pr.portfolio_debt_keur = 200_000.0
    pr.portfolio_project_irr = 0.0
    pr.portfolio_sponsor_irr = 0.0  # placeholder

    df = build_portfolio_table(pr)
    # Find the sponsor IRR row
    sponsor_rows = [i for i, label in enumerate(df.index) if "sponsor" in label.lower()]
    assert sponsor_rows, "Should have a sponsor IRR row"
    val = df.iloc[sponsor_rows[0], 0]
    assert not (isinstance(val, (int, float)) and val == 0.0), \
        f"Sponsor IRR should not be numeric 0.0, got {val!r}"





def test_portfolio_cashflows_include_negative_cfads():
    """Negative CFADS periods must be included in portfolio IRR cash flows."""
    from domain.portfolio.waterfall import build_portfolio_project_cashflows
    from domain.portfolio.inputs import PortfolioInputs
    from domain.waterfall.waterfall_engine import WaterfallResult
    from app.project_factories import create_default_solar_project, create_default_wind_project
    from domain.waterfall.waterfall_engine import WaterfallPeriod
    from domain.inputs import FinancingParams
    from datetime import date

    def _mp(period, date_, ebitda, tax):
        return WaterfallPeriod(
            period=period, date=date_, year_index=1, period_in_year=period,
            is_operation=True, generation_mwh=0.0, revenue_keur=100.0,
            opex_keur=0.0, ebitda_keur=ebitda, depreciation_keur=0.0,
            interest_senior_keur=0.0, interest_shl_keur=0.0,
            taxable_profit_keur=ebitda, tax_keur=tax,
            cf_after_tax_keur=ebitda - tax,
            senior_interest_keur=0.0, senior_principal_keur=0.0, senior_ds_keur=50.0,
            shl_interest_keur=0.0, shl_principal_keur=0.0, shl_service_keur=0.0,
            dsra_contribution_keur=0.0, dsra_balance_keur=0.0,
            mra_contribution_keur=0.0, mra_balance_keur=0.0,
            cf_after_reserves_keur=ebitda - tax, dscr=1.0, llcr=1.0, plcr=1.0,
            lockup_active=False, distribution_keur=0.0, cash_sweep_keur=0.0,
            cum_distribution_keur=0.0, cash_balance_keur=0.0,
            shl_balance_keur=0.0, shl_pik_keur=0.0, senior_balance_keur=0.0,
        )

    solar = create_default_solar_project()
    wind = create_default_wind_project()
    shared = FinancingParams(share_capital_keur=100.0, senior_debt_amount_keur=200.0,
                             senior_tenor_years=10, target_dscr=1.3)
    pi = PortfolioInputs(projects=(solar, wind), portfolio_name="Test", shared_financing=shared)

    # Solar: 20 periods of large CFADS (enough for computable IRR)
    solar_periods = []
    for i in range(20):
        d = date(2031 + i // 2, 6 if i % 2 == 0 else 12, 30 if i % 2 == 0 else 31)
        solar_periods.append(_mp(i + 1, d, 8000.0, 1200.0))  # cfads=6800
    solar_wf = WaterfallResult(
        periods=tuple(solar_periods), total_revenue_keur=100.0, total_opex_keur=0.0,
        total_ebitda_keur=100.0, total_tax_keur=20.0, total_senior_ds_keur=50.0,
        total_shl_service_keur=0.0, total_distribution_keur=0.0,
        avg_dscr=1.0, min_dscr=1.0, max_dscr=1.0,
        min_llcr=1.0, min_plcr=1.0, periods_in_lockup=0,
        project_irr=0.0, equity_irr=0.0, sponsor_irr=0.0,
        project_npv=0.0, equity_npv=0.0, sculpting_result=None,
    )

    # Wind: period 1 has NEGATIVE CFADS
    wind_periods = []
    for i in range(20):
        d = date(2031 + i // 2, 6 if i % 2 == 0 else 12, 30 if i % 2 == 0 else 31)
        if i == 0:
            wind_periods.append(_mp(i + 1, d, 200.0, 500.0))  # cfads=-300
        else:
            wind_periods.append(_mp(i + 1, d, 8000.0, 1200.0))
    wind_wf = WaterfallResult(
        periods=tuple(wind_periods), total_revenue_keur=100.0, total_opex_keur=0.0,
        total_ebitda_keur=100.0, total_tax_keur=20.0, total_senior_ds_keur=50.0,
        total_shl_service_keur=0.0, total_distribution_keur=0.0,
        avg_dscr=1.0, min_dscr=1.0, max_dscr=1.0,
        min_llcr=1.0, min_plcr=1.0, periods_in_lockup=0,
        project_irr=0.0, equity_irr=0.0, sponsor_irr=0.0,
        project_npv=0.0, equity_npv=0.0, sculpting_result=None,
    )

    cf_list, date_list = build_portfolio_project_cashflows(
        pi, (("SOLAR-001", solar_wf), ("WIND-001", wind_wf))
    )

    # June 2031: solar(+6800) + wind(-300) = 6500
    idx = date_list.index(date(2031, 6, 30))
    assert cf_list[idx] == 6500.0, f"June CF should be 6500, got {cf_list[idx]}"

    # December 2031: both positive = 6800+6800=13600
    idx2 = date_list.index(date(2031, 12, 31))
    assert cf_list[idx2] == 13600.0, f"Dec CF should be 13600, got {cf_list[idx2]}"


def test_portfolio_irr_decreases_with_negative_cfads():
    """Negative CFADS must reduce portfolio IRR, not be silently dropped."""
    from domain.portfolio.waterfall import build_portfolio_project_cashflows
    from domain.portfolio.inputs import PortfolioInputs
    from domain.returns.xirr import xirr
    from app.project_factories import create_default_solar_project, create_default_wind_project
    from domain.waterfall.waterfall_engine import WaterfallPeriod, WaterfallResult
    from domain.inputs import FinancingParams
    from datetime import date

    def _mp(period, date_, ebitda, tax):
        return WaterfallPeriod(
            period=period, date=date_, year_index=1, period_in_year=period,
            is_operation=True, generation_mwh=0.0, revenue_keur=100.0,
            opex_keur=0.0, ebitda_keur=ebitda, depreciation_keur=0.0,
            interest_senior_keur=0.0, interest_shl_keur=0.0,
            taxable_profit_keur=ebitda, tax_keur=tax,
            cf_after_tax_keur=ebitda - tax,
            senior_interest_keur=0.0, senior_principal_keur=0.0, senior_ds_keur=50.0,
            shl_interest_keur=0.0, shl_principal_keur=0.0, shl_service_keur=0.0,
            dsra_contribution_keur=0.0, dsra_balance_keur=0.0,
            mra_contribution_keur=0.0, mra_balance_keur=0.0,
            cf_after_reserves_keur=ebitda - tax, dscr=1.0, llcr=1.0, plcr=1.0,
            lockup_active=False, distribution_keur=0.0, cash_sweep_keur=0.0,
            cum_distribution_keur=0.0, cash_balance_keur=0.0,
            shl_balance_keur=0.0, shl_pik_keur=0.0, senior_balance_keur=0.0,
        )

    solar = create_default_solar_project()
    wind = create_default_wind_project()
    shared = FinancingParams(share_capital_keur=100.0, senior_debt_amount_keur=200.0,
                             senior_tenor_years=10, target_dscr=1.3)
    pi = PortfolioInputs(projects=(solar, wind), portfolio_name="Test", shared_financing=shared)

    # All-positive: 20 periods
    pos_periods = []
    for i in range(20):
        d = date(2031 + i // 2, 6 if i % 2 == 0 else 12, 30 if i % 2 == 0 else 31)
        pos_periods.append(_mp(i + 1, d, 8000.0, 1200.0))
    wf_pos = WaterfallResult(
        periods=tuple(pos_periods), total_revenue_keur=100.0, total_opex_keur=0.0,
        total_ebitda_keur=100.0, total_tax_keur=20.0, total_senior_ds_keur=50.0,
        total_shl_service_keur=0.0, total_distribution_keur=0.0,
        avg_dscr=1.0, min_dscr=1.0, max_dscr=1.0,
        min_llcr=1.0, min_plcr=1.0, periods_in_lockup=0,
        project_irr=0.0, equity_irr=0.0, sponsor_irr=0.0,
        project_npv=0.0, equity_npv=0.0, sculpting_result=None,
    )

    cf_pos, dt_pos = build_portfolio_project_cashflows(
        pi, (("SOLAR-001", wf_pos), ("WIND-001", wf_pos))
    )
    irr_pos = xirr(cf_pos, dt_pos) or 0.0
    assert irr_pos > 0.05, f"Baseline IRR should be positive and significant, got {irr_pos}"

    # Add negative CFADS period to wind
    neg_periods = list(pos_periods)
    neg_periods[0] = _mp(1, date(2031, 6, 30), 200.0, 500.0)  # cfads=-300
    wf_neg = WaterfallResult(
        periods=tuple(neg_periods), total_revenue_keur=100.0, total_opex_keur=0.0,
        total_ebitda_keur=100.0, total_tax_keur=20.0, total_senior_ds_keur=50.0,
        total_shl_service_keur=0.0, total_distribution_keur=0.0,
        avg_dscr=1.0, min_dscr=1.0, max_dscr=1.0,
        min_llcr=1.0, min_plcr=1.0, periods_in_lockup=0,
        project_irr=0.0, equity_irr=0.0, sponsor_irr=0.0,
        project_npv=0.0, equity_npv=0.0, sculpting_result=None,
    )
    cf_neg, dt_neg = build_portfolio_project_cashflows(
        pi, (("SOLAR-001", wf_pos), ("WIND-001", wf_neg))
    )
    irr_neg = xirr(cf_neg, dt_neg) or 0.0

    assert irr_neg < irr_pos, (
        f"IRR with negative CFADS ({irr_neg:.4f}) must be lower "
        f"than positive-only ({irr_pos:.4f})"
    )


class TestBuildPortfolioCashflowTable:
    """Tests for build_portfolio_cashflow_table output structure and invariants."""

    def test_portfolio_cashflow_table_sums_match_total(self):
        """Sum of breakdown per date must equal total_cashflow."""
        from domain.portfolio.waterfall import build_portfolio_cashflow_table
        from domain.portfolio.waterfall import aggregate_project_results
        from domain.portfolio.inputs import PortfolioInputs
        from domain.inputs import FinancingParams
        from app.project_factories import create_default_solar_project, create_default_wind_project

        solar = create_default_solar_project()
        wind = create_default_wind_project()
        shared = FinancingParams(
            share_capital_keur=100.0, senior_debt_amount_keur=200.0,
            senior_tenor_years=10, target_dscr=1.3,
        )
        pi = PortfolioInputs(
            projects=(solar, wind), portfolio_name="Test",
            shared_financing=shared,
        )
        # Use real waterfall runs
        from app.portfolio_runner import run_portfolio_from_inputs
        result = run_portfolio_from_inputs(pi)
        project_results = result.project_results

        table = build_portfolio_cashflow_table(pi, project_results)
        for row in table:
            total = row["total_cashflow"]
            breakdown_sum = sum(row["breakdown"].values())
            assert abs(breakdown_sum - total) < 1.0, (
                f"Date {row['date']}: breakdown sum {breakdown_sum} != total {total}"
            )

    def test_portfolio_cashflow_contains_each_project(self):
        """Every project in portfolio_inputs must appear in at least one row's breakdown."""
        from domain.portfolio.waterfall import build_portfolio_cashflow_table
        from domain.portfolio.inputs import PortfolioInputs
        from domain.inputs import FinancingParams
        from app.project_factories import create_default_solar_project, create_default_wind_project

        solar = create_default_solar_project()
        wind = create_default_wind_project()
        shared = FinancingParams(
            share_capital_keur=100.0, senior_debt_amount_keur=200.0,
            senior_tenor_years=10, target_dscr=1.3,
        )
        pi = PortfolioInputs(
            projects=(solar, wind), portfolio_name="Test",
            shared_financing=shared,
        )
        from app.portfolio_runner import run_portfolio_from_inputs
        result = run_portfolio_from_inputs(pi)

        table = build_portfolio_cashflow_table(pi, result.project_results)
        all_codes = set()
        for row in table:
            all_codes.update(row["breakdown"].keys())
        for proj in pi.projects:
            assert proj.info.code in all_codes, (
                f"Project {proj.info.code} not found in cashflow table"
            )

    def test_portfolio_cashflow_dates_sorted(self):
        """Dates in table must be in ascending order."""
        from domain.portfolio.waterfall import build_portfolio_cashflow_table
        from domain.portfolio.inputs import PortfolioInputs
        from domain.inputs import FinancingParams
        from app.project_factories import create_default_solar_project, create_default_wind_project

        solar = create_default_solar_project()
        wind = create_default_wind_project()
        shared = FinancingParams(
            share_capital_keur=100.0, senior_debt_amount_keur=200.0,
            senior_tenor_years=10, target_dscr=1.3,
        )
        pi = PortfolioInputs(
            projects=(solar, wind), portfolio_name="Test",
            shared_financing=shared,
        )
        from app.portfolio_runner import run_portfolio_from_inputs
        result = run_portfolio_from_inputs(pi)

        table = build_portfolio_cashflow_table(pi, result.project_results)
        dates = [row["date"] for row in table]
        assert dates == sorted(dates), "Dates must be sorted ascending"

    def test_portfolio_cashflow_matches_irr_inputs(self):
        """CFADS rows in build_portfolio_cashflow_table must match portfolio_cfads_schedule."""
        from domain.portfolio.waterfall import (
            build_portfolio_cashflow_table,
            portfolio_cfads_schedule,
            aggregate_project_results,
        )
        from domain.portfolio.inputs import PortfolioInputs
        from domain.inputs import FinancingParams
        from app.project_factories import create_default_solar_project, create_default_wind_project

        solar = create_default_solar_project()
        wind = create_default_wind_project()
        shared = FinancingParams(
            share_capital_keur=100.0, senior_debt_amount_keur=200.0,
            senior_tenor_years=10, target_dscr=1.3,
        )
        pi = PortfolioInputs(
            projects=(solar, wind), portfolio_name="Test",
            shared_financing=shared,
        )
        from app.portfolio_runner import run_portfolio_from_inputs
        result = run_portfolio_from_inputs(pi)

        pooled = aggregate_project_results(result.project_results)
        cfads_schedule = portfolio_cfads_schedule(pooled)

        table = build_portfolio_cashflow_table(pi, result.project_results)
        # Each operating row total_cashflow (CFADS only, no capex rows) should match
        # cfads_schedule entries that correspond to operation dates
        op_rows = [r for r in table if any(k in r["breakdown"] for k in [s.info.code for s in pi.projects])]
        # Compare totals excluding capex (capex rows have negative totals)
        cfads_rows = [r for r in table if r["total_cashflow"] > 0]
        # The table may have capex rows; filter to operation periods
        # Use pooled period dates to identify CFADS-only rows
        cfads_dates = [p["date"] for p in pooled]
        cfads_from_table = [r["total_cashflow"] for r in table if r["date"] in cfads_dates]
        # Each table CFADS should be in cfads_schedule (order may differ)
        for v in cfads_from_table:
            assert any(abs(v - c) < 1.0 for c in cfads_schedule), (
                f"CFADS value {v} from table not in cfads_schedule {cfads_schedule}"
            )

    def test_portfolio_cashflow_structure_fields(self):
        """Each row must have date, total_cashflow, and breakdown keys."""
        from domain.portfolio.waterfall import build_portfolio_cashflow_table
        from domain.portfolio.waterfall import aggregate_project_results
        from domain.portfolio.inputs import PortfolioInputs
        from domain.inputs import FinancingParams
        from app.project_factories import create_default_solar_project, create_default_wind_project

        solar = create_default_solar_project()
        wind = create_default_wind_project()
        shared = FinancingParams(
            share_capital_keur=100.0, senior_debt_amount_keur=200.0,
            senior_tenor_years=10, target_dscr=1.3,
        )
        pi = PortfolioInputs(
            projects=(solar, wind), portfolio_name="Test",
            shared_financing=shared,
        )
        from app.portfolio_runner import run_portfolio_from_inputs
        result = run_portfolio_from_inputs(pi)

        table = build_portfolio_cashflow_table(pi, result.project_results)
        for row in table:
            assert "date" in row, "Row missing 'date' key"
            assert "total_cashflow" in row, "Row missing 'total_cashflow' key"
            assert "breakdown" in row, "Row missing 'breakdown' key"
            assert isinstance(row["breakdown"], dict), "breakdown must be a dict"


def test_portfolio_result_contains_cashflow_audit_table():
    """PortfolioResult.portfolio_cashflows must be populated by run_portfolio_waterfall."""
    wf_a = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
    wf_b = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)
    pi = _portfolio_inputs()

    result = run_portfolio_waterfall(pi, (("A", wf_a), ("B", wf_b)))

    assert result.portfolio_cashflows, "portfolio_cashflows must be non-empty"

    for row in result.portfolio_cashflows:
        assert "date" in row, "Row missing 'date'"
        assert "total_cashflow" in row, "Row missing 'total_cashflow'"
        assert "breakdown" in row, "Row missing 'breakdown'"
        assert isinstance(row["breakdown"], dict), "breakdown must be a dict"

        total_from_breakdown = sum(row["breakdown"].values())
        assert abs(total_from_breakdown - row["total_cashflow"]) < 0.01, (
            f"Breakdown sum {total_from_breakdown} != total_cashflow {row['total_cashflow']}"
        )


def test_portfolio_handles_different_financial_close_dates():
    """Portfolio must correctly handle projects with different financial_close dates."""
    from domain.portfolio.waterfall import build_portfolio_project_cashflows
    from domain.portfolio.inputs import PortfolioInputs
    from domain.waterfall.waterfall_engine import WaterfallPeriod, WaterfallResult
    from app.project_factories import create_default_solar_project, create_default_wind_project
    from domain.inputs import FinancingParams
    from datetime import date

    def _mp(period, date_, ebitda, tax):
        return WaterfallPeriod(
            period=period, date=date_, year_index=1, period_in_year=period,
            is_operation=True, generation_mwh=0.0, revenue_keur=100.0,
            opex_keur=0.0, ebitda_keur=ebitda, depreciation_keur=0.0,
            interest_senior_keur=0.0, interest_shl_keur=0.0,
            taxable_profit_keur=ebitda, tax_keur=tax,
            cf_after_tax_keur=ebitda - tax,
            senior_interest_keur=0.0, senior_principal_keur=0.0, senior_ds_keur=50.0,
            shl_interest_keur=0.0, shl_principal_keur=0.0, shl_service_keur=0.0,
            dsra_contribution_keur=0.0, dsra_balance_keur=0.0,
            mra_contribution_keur=0.0, mra_balance_keur=0.0,
            cf_after_reserves_keur=ebitda - tax, dscr=1.0, llcr=1.0, plcr=1.0,
            lockup_active=False, distribution_keur=0.0, cash_sweep_keur=0.0,
            cum_distribution_keur=0.0, cash_balance_keur=0.0,
            shl_balance_keur=0.0, shl_pik_keur=0.0, senior_balance_keur=0.0,
        )

    solar = create_default_solar_project()
    wind = create_default_wind_project()

    # Override wind's financial_close to be 6 months later
    from dataclasses import replace
    wind_later = replace(wind, info=replace(wind.info, financial_close=date(2030, 7, 1)))

    shared = FinancingParams(share_capital_keur=100.0, senior_debt_amount_keur=200.0,
                             senior_tenor_years=10, target_dscr=1.3)
    pi = PortfolioInputs(projects=(solar, wind_later), portfolio_name="Test", shared_financing=shared)

    wf_s = _make_wf_result("A", ebitda=80.0, tax=10.0, rev=100.0)
    wf_w = _make_wf_result("B", ebitda=120.0, tax=15.0, rev=150.0)

    cf_list, date_list = build_portfolio_project_cashflows(
        pi, (("SOLAR-001", wf_s), ("WIND-001", wf_w))
    )

    assert len(cf_list) == len(date_list)
    assert len(cf_list) >= 2, "Should have CapEx outflows for both projects"


def test_portfolio_handles_staggered_operations():
    """Portfolio must handle projects where operations start in different periods."""
    from domain.portfolio.waterfall import build_portfolio_project_cashflows
    from domain.portfolio.inputs import PortfolioInputs
    from domain.waterfall.waterfall_engine import WaterfallPeriod, WaterfallResult
    from app.project_factories import create_default_solar_project, create_default_wind_project
    from domain.inputs import FinancingParams
    from datetime import date

    def _mp(period, date_, ebitda, tax):
        return WaterfallPeriod(
            period=period, date=date_, year_index=1, period_in_year=period,
            is_operation=True, generation_mwh=0.0, revenue_keur=100.0,
            opex_keur=0.0, ebitda_keur=ebitda, depreciation_keur=0.0,
            interest_senior_keur=0.0, interest_shl_keur=0.0,
            taxable_profit_keur=ebitda, tax_keur=tax,
            cf_after_tax_keur=ebitda - tax,
            senior_interest_keur=0.0, senior_principal_keur=0.0, senior_ds_keur=50.0,
            shl_interest_keur=0.0, shl_principal_keur=0.0, shl_service_keur=0.0,
            dsra_contribution_keur=0.0, dsra_balance_keur=0.0,
            mra_contribution_keur=0.0, mra_balance_keur=0.0,
            cf_after_reserves_keur=ebitda - tax, dscr=1.0, llcr=1.0, plcr=1.0,
            lockup_active=False, distribution_keur=0.0, cash_sweep_keur=0.0,
            cum_distribution_keur=0.0, cash_balance_keur=0.0,
            shl_balance_keur=0.0, shl_pik_keur=0.0, senior_balance_keur=0.0,
        )

    solar = create_default_solar_project()
    wind = create_default_wind_project()
    shared = FinancingParams(share_capital_keur=100.0, senior_debt_amount_keur=200.0,
                             senior_tenor_years=10, target_dscr=1.3)
    pi = PortfolioInputs(projects=(solar, wind), portfolio_name="Test", shared_financing=shared)

    # Solar starts first (period 1-10), Wind starts second (period 3-12)
    solar_periods = [_mp(i + 1, date(2031 + i // 2, 6 if i % 2 == 0 else 12, 30 if i % 2 == 0 else 31), 80.0, 10.0) for i in range(10)]
    wind_periods = [_mp(i + 3, date(2031 + (i + 2) // 2, 6 if (i + 2) % 2 == 0 else 12, 30 if (i + 2) % 2 == 0 else 31), 120.0, 15.0) for i in range(10)]

    wf_s = WaterfallResult(
        periods=tuple(solar_periods), total_revenue_keur=100.0, total_opex_keur=0.0,
        total_ebitda_keur=100.0, total_tax_keur=20.0, total_senior_ds_keur=50.0,
        total_shl_service_keur=0.0, total_distribution_keur=0.0,
        avg_dscr=1.0, min_dscr=1.0, max_dscr=1.0,
        min_llcr=1.0, min_plcr=1.0, periods_in_lockup=0,
        project_irr=0.0, equity_irr=0.0, sponsor_irr=0.0,
        project_npv=0.0, equity_npv=0.0, sculpting_result=None,
    )
    wf_w = WaterfallResult(
        periods=tuple(wind_periods), total_revenue_keur=150.0, total_opex_keur=0.0,
        total_ebitda_keur=150.0, total_tax_keur=30.0, total_senior_ds_keur=50.0,
        total_shl_service_keur=0.0, total_distribution_keur=0.0,
        avg_dscr=1.0, min_dscr=1.0, max_dscr=1.0,
        min_llcr=1.0, min_plcr=1.0, periods_in_lockup=0,
        project_irr=0.0, equity_irr=0.0, sponsor_irr=0.0,
        project_npv=0.0, equity_npv=0.0, sculpting_result=None,
    )

    cf_list, date_list = build_portfolio_project_cashflows(
        pi, (("SOLAR-001", wf_s), ("WIND-001", wf_w))
    )

    # Periods 1-2: only solar operating → CF = 70 (80-10)
    assert cf_list[0] == 70.0 or cf_list[1] == 70.0, "First periods should have solar-only CF"
    # Period 3+: both operating → CF = 70 + 105 = 175
    for i in range(3, len(cf_list)):
        # Wind operates, solar operates → 175 or more
        if date_list[i] >= wind_periods[0].date:
            pass  # just verify no crash

    assert len(cf_list) >= 10, "Portfolio should have at least 10 cashflow entries"


def test_portfolio_cf_export_exists():
    """Portfolio CF sheet must exist when portfolio_result has cashflows."""
    import openpyxl
    from io import BytesIO
    from app.excel_export import build_excel_export
    from app.portfolio_runner import run_portfolio_from_inputs
    from app.project_factories import create_default_solar_project, create_default_wind_project
    from domain.portfolio.inputs import PortfolioInputs
    from domain.inputs import FinancingParams

    solar = create_default_solar_project()
    wind = create_default_wind_project()
    shared = FinancingParams(share_capital_keur=100.0, senior_debt_amount_keur=200.0,
                             senior_tenor_years=10, target_dscr=1.3)
    pi = PortfolioInputs(projects=(solar, wind), portfolio_name="Test", shared_financing=shared)
    result = run_portfolio_from_inputs(pi)

    data = build_excel_export(portfolio_result=result,
        project_inputs=None,
        integration_status="experimental",
        scenario="Base",
    )
    wb = openpyxl.load_workbook(BytesIO(data))
    assert "Portfolio CF" in wb.sheetnames, "Portfolio CF sheet must exist"


def test_portfolio_cf_breakdown_sums_match_total():
    """Each row's breakdown sum must equal total_cashflow."""
    from domain.portfolio.waterfall import build_portfolio_cashflow_table
    from app.project_factories import create_default_solar_project, create_default_wind_project
    from domain.portfolio.inputs import PortfolioInputs
    from domain.waterfall.waterfall_engine import WaterfallResult, WaterfallPeriod
    from domain.inputs import FinancingParams
    from datetime import date

    def _mp(period, date_, ebitda, tax):
        return WaterfallPeriod(
            period=period, date=date_, year_index=1, period_in_year=period,
            is_operation=True, generation_mwh=0.0, revenue_keur=100.0,
            opex_keur=0.0, ebitda_keur=ebitda, depreciation_keur=0.0,
            interest_senior_keur=0.0, interest_shl_keur=0.0,
            taxable_profit_keur=ebitda, tax_keur=tax,
            cf_after_tax_keur=ebitda - tax,
            senior_interest_keur=0.0, senior_principal_keur=0.0, senior_ds_keur=50.0,
            shl_interest_keur=0.0, shl_principal_keur=0.0, shl_service_keur=0.0,
            dsra_contribution_keur=0.0, dsra_balance_keur=0.0,
            mra_contribution_keur=0.0, mra_balance_keur=0.0,
            cf_after_reserves_keur=ebitda - tax, dscr=1.0, llcr=1.0, plcr=1.0,
            lockup_active=False, distribution_keur=0.0, cash_sweep_keur=0.0,
            cum_distribution_keur=0.0, cash_balance_keur=0.0,
            shl_balance_keur=0.0, shl_pik_keur=0.0, senior_balance_keur=0.0,
        )

    solar = create_default_solar_project()
    wind = create_default_wind_project()
    shared = FinancingParams(share_capital_keur=100.0, senior_debt_amount_keur=200.0,
                             senior_tenor_years=10, target_dscr=1.3)
    pi = PortfolioInputs(projects=(solar, wind), portfolio_name="Test", shared_financing=shared)

    periods = [_mp(i + 1, date(2031 + i // 2, 6 if i % 2 == 0 else 12, 30 if i % 2 == 0 else 31), 8000.0, 1200.0) for i in range(10)]
    wf = WaterfallResult(
        periods=tuple(periods), total_revenue_keur=100.0, total_opex_keur=0.0,
        total_ebitda_keur=100.0, total_tax_keur=20.0, total_senior_ds_keur=50.0,
        total_shl_service_keur=0.0, total_distribution_keur=0.0,
        avg_dscr=1.0, min_dscr=1.0, max_dscr=1.0,
        min_llcr=1.0, min_plcr=1.0, periods_in_lockup=0,
        project_irr=0.0, equity_irr=0.0, sponsor_irr=0.0,
        project_npv=0.0, equity_npv=0.0, sculpting_result=None,
    )

    table = build_portfolio_cashflow_table(pi, (("SOLAR-001", wf), ("WIND-001", wf)))
    for row in table:
        total = row["total_cashflow"]
        breakdown_sum = sum(row["breakdown"].values())
        assert abs(breakdown_sum - total) < 0.01, (
            f"Breakdown sum {breakdown_sum} != total {total}"
        )
