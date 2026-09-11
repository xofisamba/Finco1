"""R5 / F04 — Snapshot-authoritative institutional export.

Closes audit finding F04: the institutional XLSX and runtime-summary CSV
export routes reconstructed the project from ``PROJECT_FACTORIES`` and ran
that, ignoring the user's persisted working copy and selected scenario —
exporting factory economics (e.g. TUHO 35 MW @ 18 % CIT) for a saved working
copy with materially different edits.

R5 authority chain (identical to Workbook V2 Run):

    persisted draft snapshot
        → ProjectInputSet / to_projectinputs()
        → active-scenario CAPEX replace-fold + OPEX additive-fold
        → canonical execute_production_waterfall
        → export bundle / workbook bytes

The factory remains the authority ONLY for factory-template references and
for context-less legacy callers.  ``financial_engine/`` and ``finco_core/``
are untouched.
"""
from __future__ import annotations

import os
import re
import urllib.parse

os.environ.setdefault("FINCO_WORKBOOK_V2", "1")
os.environ.setdefault("FINCO_SECRET_KEY", "test-secret-r5")

import openpyxl  # noqa: E402
import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import main_web  # noqa: E402
from app.auth import COOKIE_NAME, create_session_token, decode_session_token  # noqa: E402

CIT_FIELD = "tax.assumptions.cit_rate_pct"
EPC_FIELD = "capex.C.epc_contract"


def _client() -> TestClient:
    tc = TestClient(main_web.app, follow_redirects=False)
    tc.cookies.set(COOKIE_NAME, _SESSION)
    return tc


_SESSION = create_session_token()
_USER_ID = decode_session_token(_SESSION).user_id


def _create(client: TestClient, name: str, template_source: str, **overrides) -> str:
    data = {
        "project_name": name,
        "project_type": "Wind" if template_source in ("tuho", "generic_wind") else "Solar",
        "template_source": template_source,
        "country_market": "Croatia" if template_source == "tuho" else "Poland",
        "capacity_mw": "53",
        "cod_date": "2029-12-30" if template_source == "tuho" else "2028-06-01",
        "construction_months": "18" if template_source == "tuho" else "12",
        "horizon_years": "30",
        "tariff_eur_mwh": "60" if template_source == "tuho" else "62",
        "ppa_term_years": "20",
        "p50_hours": "4164" if template_source == "tuho" else "1900",
        "opex_y1_keur": "1998.01" if template_source == "tuho" else "850",
        "total_capex_keur": "70691.53944444444" if template_source == "tuho" else "58000",
        "gearing_pct": "",
        "interest_rate_pct": "5.98" if template_source == "tuho" else "4.5",
        "tenor_years": "14" if template_source == "tuho" else "18",
        "target_dscr": "1.2" if template_source == "tuho" else "1.25",
        **overrides,
    }
    resp = client.post("/projects/create", data=data, follow_redirects=False)
    redirect = resp.headers.get("hx-redirect") or resp.headers.get("location", "")
    parsed = urllib.parse.parse_qs(urllib.parse.urlparse(redirect).query)
    code = parsed.get("project", [None])[0]
    assert code, f"create failed: {resp.status_code} {resp.text[:200]}"
    return code


def _hash(client: TestClient, code: str) -> tuple[str, str]:
    body = client.get(f"/v2/workbook?project={code}").text
    return (
        re.search(r'data-content-hash="([^"]+)"', body).group(1),
        re.search(r'data-workbook-version="([^"]+)"', body).group(1),
    )


def _edit(client: TestClient, code: str, field_id: str, value: str, sheet: str) -> None:
    ch, wv = _hash(client, code)
    resp = client.post(
        "/v2/workbook/update",
        data={"field_id": field_id, "value": value, "project": code,
              "workbook_version": wv, "content_hash": ch, "sheet_id": sheet},
        headers={"HX-Request": "true"},
    )
    assert resp.status_code == 200, f"edit {field_id}: {resp.status_code} {resp.text[:200]}"


def _run(client: TestClient, code: str) -> None:
    ch, wv = _hash(client, code)
    resp = client.post(
        "/v2/workbook/run",
        data={"project": code, "content_hash": ch, "workbook_version": wv},
        headers={"HX-Request": "true"}, follow_redirects=False,
    )
    assert resp.status_code == 200, f"run failed: {resp.status_code}"


def _export_workbook(client: TestClient, code: str) -> openpyxl.Workbook:
    resp = client.get(f"/exports/institutional-workbook.xlsx?project={code}")
    assert resp.status_code == 200, f"export failed: {resp.status_code} {resp.text[:200]}"
    import io

    return openpyxl.load_workbook(io.BytesIO(resp.content))


def _export_csv(client: TestClient, code: str) -> str:
    resp = client.get(f"/exports/runtime-summary.csv?project={code}")
    assert resp.status_code == 200, f"csv export failed: {resp.status_code}"
    return resp.content.decode("utf-8")


def _sheet_values(wb: openpyxl.Workbook, title: str) -> dict[str, object]:
    """Flatten a worksheet to {first-column label: second-column value}."""
    ws = wb[title]
    out: dict[str, object] = {}
    for row in ws.iter_rows():
        cells = [c.value for c in row]
        if cells and cells[0] is not None:
            out[str(cells[0])] = cells[1] if len(cells) > 1 else None
    return out


def _runtime_summary(wb: openpyxl.Workbook) -> dict[str, object]:
    return _sheet_values(wb, "Runtime Summary")


def _force_active_scenario_id(project_id: str, scenario_id: str | None) -> None:
    """Overwrite the workspace's active_scenario_id bypassing normal select flow.
    Used to inject invalid scenario IDs for fail-closed tests."""
    from app.persistence.workspace_repository import get_workspace_state, save_workspace_state

    ws = get_workspace_state(_USER_ID, project_id)
    assert ws is not None
    save_workspace_state(
        user_id=_USER_ID,
        project_id=project_id,
        project_code=ws.project_code,
        draft_snapshot=dict(ws.draft_snapshot),
        saved_snapshot=dict(ws.saved_snapshot),
        governance_state=dict(ws.governance_state or {}),
        active_scenario_id=scenario_id,
        active_scenario_name=ws.active_scenario_name if scenario_id is not None else None,
        last_runtime_snapshot=dict(ws.last_runtime_snapshot or {}),
        last_runtime_summary=dict(ws.last_runtime_summary or {}),
        last_runtime_snapshot_id=ws.last_runtime_snapshot_id,
        last_runtime_origin=ws.last_runtime_origin,
        last_runtime_scenario_id=ws.last_runtime_scenario_id,
        replay_metadata=dict(ws.replay_metadata or {}),
    )


def _saved_inputs(code: str):
    from app.persistence.projects_repository import get_project_by_code
    from app.services.export_service import (
        resolve_snapshot_authoritative_project_inputs,
    )

    rec = get_project_by_code(_USER_ID, code)
    return resolve_snapshot_authoritative_project_inputs(rec, _USER_ID)


def _canonical_run(code: str):
    """Canonical production run of the saved working copy -> (kpis, inputs)."""
    from app.api.project_runner import run_project

    pi = _saved_inputs(code)
    result = run_project("Wind", "Base", project_inputs_override=pi)
    return result["kpis"], pi


# Shared, expensive fixtures — one TUHO working copy with CIT 15 % + EPC 20000
# persisted through the REAL edit/persistence path, then Run, then export.


@pytest.fixture(scope="module")
def tuho_code() -> str:
    client = _client()
    code = _create(client, "R5 TUHO WC", "tuho")
    _edit(client, code, CIT_FIELD, "15", "tax")
    _edit(client, code, EPC_FIELD, "20000", "capex")
    _run(client, code)
    return code


@pytest.fixture(scope="module")
def tuho_wb(tuho_code) -> openpyxl.Workbook:
    return _export_workbook(_client(), tuho_code)


class TestTuhoCitAuthority:
    def test_a_export_uses_saved_15_percent_not_factory_18(self, tuho_code, tuho_wb):
        """§7-A: the export runs the saved working copy — its runtime
        economics equal a canonical run of the SAME saved inputs, and the
        saved 15 % CIT override is materially visible (the factory TUHO runs
        18 % and a different capacity, so a factory substitution cannot
        reconcile)."""
        result, saved = _canonical_run(tuho_code)
        rs = _runtime_summary(tuho_wb)
        assert saved.tax.corporate_rate_override == pytest.approx(0.15)
        assert saved.tax.country_tax_policy_id == "HR-approved-source-model-2026-v1"
        assert saved.technical.capacity_mw == pytest.approx(53.0)
        assert rs["Project IRR"] == pytest.approx(result["project_irr"], rel=1e-9)
        assert rs["Total revenue"] == pytest.approx(
            result["total_revenue_keur"], rel=1e-9)

    def test_a2_workbook_context_is_the_working_copy(self, tuho_code, tuho_wb):
        cover = _sheet_values(tuho_wb, "Cover")
        inputs = _sheet_values(tuho_wb, "Inputs")
        assert "R5 TUHO WC" in str(cover.get("Project", ""))
        # saved capacity 53 — not the factory 35
        assert inputs.get("Capacity MW") == 53

    def test_a3_capex_edit_survives_export(self, tuho_code, tuho_wb):
        """§7-B (part 1): the R1 scalar CAPEX edit is exported."""
        ws = tuho_wb["CAPEX"]
        epc_values = [
            c.value for row in ws.iter_rows() for c in row
            if c.value == "EPC Contract"
        ]
        assert epc_values, "EPC Contract row missing"
        # the edited 20000 appears on the CAPEX sheet
        all_values = [
            str(c.value) for row in ws.iter_rows() for c in row
            if c.value is not None
        ]
        assert any(v.startswith("20000") for v in all_values), all_values[:30]


class TestCombinedAuthorities:
    def test_b_persistence_run_export_keep_both_edits(self, tuho_code):
        """§7-B (part 2) + §7-G: fresh client session (close/reload) → the
        persisted edits survive reload → Run → export; BOTH the CIT override
        and the CAPEX scalar remain authoritative."""
        fresh = _client()  # new HTTP session, same persisted user state
        body = fresh.get(f"/v2/workbook?project={tuho_code}").text
        assert 'data-content-hash' in body  # reload OK
        result, saved = _canonical_run(tuho_code)
        assert saved.tax.corporate_rate_override == pytest.approx(0.15)
        assert [i.amount_keur for i in saved.capex.capex_items()][0] == 20000.0
        wb = _export_workbook(fresh, tuho_code)
        rs = _runtime_summary(wb)
        assert rs["Project IRR"] == pytest.approx(result["project_irr"], rel=1e-9)


class TestExplicitZero:
    def test_c_zero_cit_survives_export_as_zero(self):
        """§7-C: 0 is economically meaningful — a saved explicit 0 % CIT must
        export zero-tax economics, not the factory default."""
        client = _client()
        code = _create(client, "R5 TUHO Zero", "tuho")
        _edit(client, code, CIT_FIELD, "0", "tax")
        _run(client, code)
        result, saved = _canonical_run(code)
        assert saved.tax.corporate_rate_override == 0.0
        assert result["total_tax_keur"] == pytest.approx(0.0)
        wb = _export_workbook(client, code)
        rs = _runtime_summary(wb)
        # zero-CIT economics are unique — export reconciling with THIS run
        # (and not the 18% factory run) proves the explicit zero survived.
        assert rs["Project IRR"] == pytest.approx(result["project_irr"], rel=1e-9)
        assert rs["Total EBITDA"] == pytest.approx(result["total_ebitda_keur"], rel=1e-9)


class TestGenericProjects:
    @pytest.mark.parametrize("template,tech,name", [
        ("generic_solar", "Solar", "R5 Solar WC"),
        ("generic_wind", "Wind", "R5 Wind WC"),
    ])
    def test_d_e_generic_saved_state_authority(self, template, tech, name):
        """§7-D/§7-E: generic projects export their persisted edits."""
        client = _client()
        code = _create(client, name, template)
        _edit(client, code, CIT_FIELD, "28", "tax")
        _edit(client, code, EPC_FIELD, "12345", "capex")
        _run(client, code)
        result, saved = _canonical_run(code)
        assert saved.tax.corporate_rate == pytest.approx(0.28)
        assert saved.tax.corporate_rate_override is None  # no policy — legacy authority
        wb = _export_workbook(client, code)
        rs = _runtime_summary(wb)
        assert rs["Project IRR"] == pytest.approx(result["project_irr"], rel=1e-9)
        ws = wb["CAPEX"]
        values = [str(c.value) for row in ws.iter_rows() for c in row
                  if c.value is not None]
        assert any(v.startswith("12345") for v in values)


class TestOborovo:
    def test_f_oborovo_saved_state_preserved(self):
        """§7-F: an Oborovo working copy exports its persisted state via the
        same snapshot authority; its legacy tax classification is unchanged
        (no typed policy fabricated)."""
        client = _client()
        code = _create(client, "R5 Oborovo WC", "oborovo",
                       project_type="Solar", tenor_years="14", target_dscr="1.15")
        _edit(client, code, CIT_FIELD, "12", "tax")
        _run(client, code)
        result, saved = _canonical_run(code)
        assert saved.tax.country_tax_policy_id is None
        assert saved.tax.corporate_rate == pytest.approx(0.12)
        assert saved.tax.corporate_rate_override is None
        wb = _export_workbook(client, code)
        rs = _runtime_summary(wb)
        assert rs["Project IRR"] == pytest.approx(result["project_irr"], rel=1e-9)


class TestScenarioAuthority:
    def test_h_base_scenario_base_no_stale_state(self, tuho_code):
        """§7-H: Base export → Scenario export (CAPEX sub-line override) →
        Base export again. The scenario export must differ (fold applied),
        and returning to Base must reproduce the original Base economics
        with no stale overlay."""
        from app.persistence.projects_repository import get_project_by_code
        from app.persistence.scenarios_repository import (
            add_scenario, get_scenario, select_scenario,
        )
        from app.workbook.input_set import ProjectInputSet
        from app.persistence.workspace_repository import get_workspace_state

        client = _client()
        base_wb = _export_workbook(client, tuho_code)
        base_rs = _runtime_summary(base_wb)

        rec = get_project_by_code(_USER_ID, tuho_code)
        ws = get_workspace_state(_USER_ID, rec.project_id)
        base_pis = ProjectInputSet.from_snapshot(dict(ws.draft_snapshot))
        scen_name = "R5 Scenario A"
        sc = add_scenario(
            user_id=_USER_ID, project_id=rec.project_id,
            project_code=rec.project_code, scenario_name=scen_name,
            parent_scenario_id=ws.active_scenario_id or "",
            base_input_set=dict(ws.draft_snapshot),
        )
        # CAPEX sub-line override: double the first sub-line amount
        from app.services.capex_sub_lines_integration import (
            _load_active_sub_lines,
        )
        sub_lines = _load_active_sub_lines(rec.project_id)
        target = sub_lines[0]
        overrides = {
            "_capex_sub_line_overrides": {
                target.sub_line_id: float(target.amount_keur) * 2.0,
            },
        }
        update_scenario_overrides(_USER_ID, sc.scenario_id, overrides)
        assert select_scenario(_USER_ID, rec.project_id, sc.scenario_id)

        scen_wb = _export_workbook(client, tuho_code)
        scen_rs = _runtime_summary(scen_wb)
        assert scen_rs["Project IRR"] != pytest.approx(
            base_rs["Project IRR"], rel=1e-9)
        # provenance reflects the scenario
        assert "R5 Scenario A" in str(scen_rs.get("Scenario", ""))

        # back to Base: original economics restored, no stale overlay
        base_rec = get_scenario(
            scenario_id=ws.active_scenario_id, user_id=_USER_ID) if ws.active_scenario_id else None
        if base_rec is not None:
            assert select_scenario(_USER_ID, rec.project_id, base_rec.scenario_id)
        else:
            # fall back: clear the active scenario by selecting the base case
            from app.persistence.scenarios_repository import list_scenarios
            base_candidates = [
                s for s in list_scenarios(_USER_ID, rec.project_id)
                if s.is_base
            ]
            assert base_candidates, "no base scenario record"
            assert select_scenario(_USER_ID, rec.project_id, base_candidates[0].scenario_id)

        restored_wb = _export_workbook(client, tuho_code)
        restored_rs = _runtime_summary(restored_wb)
        assert restored_rs["Project IRR"] == pytest.approx(
            base_rs["Project IRR"], rel=1e-9)
        assert restored_rs["Total EBITDA"] == pytest.approx(
            base_rs["Total EBITDA"], rel=1e-9)


class TestNegativeFactoryFallback:
    def test_i_export_of_working_copy_never_touches_factories(self, tuho_code, monkeypatch):
        """§7-I: with every factory constructor instrumented to raise, a
        normal institutional export of the persisted working copy still
        succeeds — proving the export path never reconstructs economics from
        a factory for a user project."""
        from app.export import institutional_workbook as iw
        from app.export import runtime_summary as rs_mod

        def _boom(*a, **k):
            raise AssertionError("FACTORY CALLED DURING EXPORT — F04 REGRESSION")

        factory_map = {k: _boom for k in rs_mod.PROJECT_FACTORIES}
        monkeypatch.setattr(iw, "PROJECT_FACTORIES", factory_map)
        monkeypatch.setattr(rs_mod, "PROJECT_FACTORIES", factory_map)
        wb = _export_workbook(_client(), tuho_code)
        assert "Runtime Summary" in wb.sheetnames


class TestFinancialStatementAvailability:
    def test_j_export_fs_matches_canonical_fail_closed_contract(self, tuho_wb):
        """§7-J: the project runs the clean G2C production authority, so the
        legacy-statement surfaces stay honestly NOT_AVAILABLE (typed marker)
        — exactly like the canonical run's statement availability. No
        fabricated statements, no factory-economics statements."""
        tax_rows = _sheet_values(tuho_wb, "Tax")
        assert any(
            "NOT_AVAILABLE" in str(v) for v in tax_rows.values()
            if v is not None
        ) or "NOT_AVAILABLE" in [str(c.value) for row in tuho_wb["Tax"].iter_rows()
                                 for c in row if c.value is not None]
        # explicit no-legacy-fallback marker present
        text = "\n".join(
            str(c.value) for row in tuho_wb["Tax"].iter_rows()
            for c in row if c.value is not None
        )
        assert "PR8_NOT_AVAILABLE" in text
        assert "no legacy fallback" in text


class TestCsvExportAuthority:
    def test_csv_export_also_snapshot_authoritative(self, tuho_code):
        """The runtime-summary CSV export shares the F04 root cause — it must
        reconcile with the canonical run of the saved working copy too."""
        import csv
        import io

        result, _ = _canonical_run(tuho_code)
        text = _export_csv(_client(), tuho_code)
        rows = list(csv.reader(io.StringIO(text)))
        # long format: project,metric,value,...
        data = {r[1]: r[2] for r in rows[1:] if len(r) >= 3}
        assert data, text[:200]
        assert "project_irr" in data
        assert data["active_project"] != "TUHO Wind 1", (
            "CSV export reconstructed the factory project — F04 regression")
        assert float(data["project_irr"]) == pytest.approx(
            result["project_irr"], rel=1e-9)


# ── R5/F04 Correction A: post-creation scalar edits reach XLSX context ────────

CAPACITY_FIELD = "project_setup.technical.capacity_mw"
TARIFF_FIELD = "revenue.ppa.base_tariff"


class TestScalarEditReachesXlsxContext:
    """R5/F04-A: the institutional workbook context must consume the SAME
    persisted draft authority that produced the effective ProjectInputs.
    A post-creation edit to a scalar field (capacity, tariff) must appear
    in both the Inputs sheet and the runtime economics of the XLSX export."""

    def test_k_capacity_edit_reaches_inputs_sheet(self):
        """Create project at 53 MW, edit capacity to 60 MW via V2 pipeline.
        The Inputs sheet of the XLSX must show 60 MW, not 53 MW (baseline)."""
        client = _client()
        code = _create(client, "R5A Capacity Edit", "tuho")
        _run(client, code)
        wb_before = _export_workbook(client, code)
        inputs_before = _sheet_values(wb_before, "Inputs")
        assert inputs_before.get("Capacity MW") == 53, (
            f"Expected baseline 53, got {inputs_before.get('Capacity MW')}")

        # Edit capacity AFTER project creation
        _edit(client, code, CAPACITY_FIELD, "60", "inputs")
        _run(client, code)
        wb_after = _export_workbook(client, code)
        inputs_after = _sheet_values(wb_after, "Inputs")

        assert inputs_after.get("Capacity MW") == 60, (
            f"Inputs sheet still shows baseline value "
            f"({inputs_after.get('Capacity MW')}) — Gap A not closed")

        # Runtime economics must also reflect the edited capacity (not baseline)
        result_after, pi_after = _canonical_run(code)
        assert pi_after.technical.capacity_mw == pytest.approx(60.0), (
            "canonical run does not see the edited capacity")
        rs = _runtime_summary(wb_after)
        assert rs["Project IRR"] == pytest.approx(result_after["project_irr"], rel=1e-9)

    def test_l_tariff_edit_reaches_revenue_sheet(self):
        """Create project at 60 EUR/MWh, edit tariff to 75 EUR/MWh.
        Both XLSX context (Revenue sheet) and runtime economics must reflect 75, not baseline."""
        client = _client()
        code = _create(client, "R5A Tariff Edit", "tuho")
        _run(client, code)

        # Record baseline tariff visible in XLSX Revenue sheet
        wb_before = _export_workbook(client, code)
        revenue_before = _sheet_values(wb_before, "Revenue")
        # TUHO created at tariff_eur_mwh=60, expect 60 before edit
        assert revenue_before.get("PPA tariff EUR/MWh") == pytest.approx(60.0, abs=1.0), (
            f"Unexpected baseline tariff: {revenue_before.get('PPA tariff EUR/MWh')}")

        # Edit tariff to 75 AFTER project creation
        _edit(client, code, TARIFF_FIELD, "75", "inputs")
        _run(client, code)
        wb_after = _export_workbook(client, code)
        revenue_after = _sheet_values(wb_after, "Revenue")

        assert revenue_after.get("PPA tariff EUR/MWh") == pytest.approx(75.0, abs=1.0), (
            f"Revenue sheet still shows baseline tariff "
            f"({revenue_after.get('PPA tariff EUR/MWh')}) — Gap A not closed")

        result_after, pi_after = _canonical_run(code)
        assert pi_after.revenue.ppa_base_tariff == pytest.approx(75.0), (
            "canonical run does not see edited tariff")
        rs = _runtime_summary(wb_after)
        assert rs["Project IRR"] == pytest.approx(result_after["project_irr"], rel=1e-9)

    def test_m_baseline_value_differs_from_edited_draft(self):
        """Prove baseline_snapshot != edited draft for the tariff field,
        establishing the defect condition this correction closes."""
        from app.persistence.projects_repository import get_project_by_code
        from app.persistence.workspace_repository import get_workspace_state

        client = _client()
        code = _create(client, "R5A Baseline Proof", "tuho")
        # Edit tariff to 80 after creation (baseline stays at 60)
        _edit(client, code, TARIFF_FIELD, "80", "inputs")
        _run(client, code)

        rec = get_project_by_code(_USER_ID, code)
        ws = get_workspace_state(_USER_ID, rec.project_id)
        baseline = dict(rec.baseline_snapshot or {})
        draft_tariff_key = "rev_ppa_base_tariff"
        baseline_tariff = float(baseline.get("tariff_eur_mwh") or baseline.get("rev_ppa_base_tariff") or 60)
        draft_tariff = float(ws.draft_snapshot.get(draft_tariff_key) or ws.draft_snapshot.get("tariff_eur_mwh") or 0)

        # This proves the defect condition (baseline ≠ draft)
        assert baseline_tariff != pytest.approx(80.0), (
            "baseline_snapshot already holds the edited value — test setup error")
        assert draft_tariff == pytest.approx(80.0, abs=2.0), (
            f"draft_snapshot does not hold edited tariff: {draft_tariff}")

        # Correction A must close this: Revenue sheet shows 80, not baseline
        wb = _export_workbook(client, code)
        revenue = _sheet_values(wb, "Revenue")
        tariff_shown = revenue.get("PPA tariff EUR/MWh")
        assert tariff_shown == pytest.approx(80.0, abs=1.0), (
            f"XLSX Revenue sheet shows baseline tariff ({tariff_shown}) instead of edited draft (80) — Gap A not closed")

    def test_n_close_reload_export_preserves_scalar_edits(self):
        """A fresh client session (close/reload) must still see the edited
        tariff in the XLSX — proving that persistence survives the round-trip."""
        client = _client()
        code = _create(client, "R5A Round-trip", "tuho")
        _edit(client, code, TARIFF_FIELD, "72", "inputs")
        _run(client, code)

        fresh = _client()  # new HTTP session
        wb = _export_workbook(fresh, code)
        revenue = _sheet_values(wb, "Revenue")
        assert revenue.get("PPA tariff EUR/MWh") == pytest.approx(72.0, abs=1.0), (
            f"Reloaded XLSX Revenue sheet shows stale tariff: {revenue.get('PPA tariff EUR/MWh')}")


# ── R5/F04 Correction B: scenario fail-closed matrix ─────────────────────────

class TestScenarioFailClosedMatrix:
    """R5/F04-B: export must mirror Workbook V2 Run's scenario validation.
    When active_scenario_id is set to an invalid scenario the export must
    fail closed with an HTTP error, never silently substitute Base economics."""

    def _setup_project_with_active_scenario(self, client: TestClient, name: str):
        """Create a project, add a scenario, make it active, return (code, scenario_id)."""
        from app.persistence.projects_repository import get_project_by_code
        from app.persistence.scenarios_repository import add_scenario, select_scenario
        from app.persistence.workspace_repository import get_workspace_state

        code = _create(client, name, "tuho")
        _edit(client, code, CIT_FIELD, "15", "tax")
        _run(client, code)
        rec = get_project_by_code(_USER_ID, code)
        ws = get_workspace_state(_USER_ID, rec.project_id)
        sc = add_scenario(
            user_id=_USER_ID, project_id=rec.project_id,
            project_code=rec.project_code, scenario_name=f"{name} Sc",
            parent_scenario_id=ws.active_scenario_id or "",
            base_input_set=dict(ws.draft_snapshot),
        )
        assert select_scenario(_USER_ID, rec.project_id, sc.scenario_id)
        return code, sc.scenario_id

    def test_o_valid_active_scenario_export_succeeds(self):
        """A valid active scenario must not cause the export to fail.
        Both CSV and XLSX exports must return 200 (not 400) when a valid
        scenario is active — the scenario path is exercised without error."""
        from app.persistence.projects_repository import get_project_by_code
        from app.persistence.scenarios_repository import add_scenario, select_scenario
        from app.persistence.workspace_repository import get_workspace_state

        client = _client()
        code = _create(client, "R5B Valid Scenario", "tuho")
        _run(client, code)

        rec = get_project_by_code(_USER_ID, code)
        ws = get_workspace_state(_USER_ID, rec.project_id)
        sc = add_scenario(
            user_id=_USER_ID, project_id=rec.project_id,
            project_code=rec.project_code, scenario_name="R5B Valid Sc",
            parent_scenario_id=ws.active_scenario_id or "",
            base_input_set=dict(ws.draft_snapshot),
        )
        assert select_scenario(_USER_ID, rec.project_id, sc.scenario_id)

        # Both exports must succeed (200) with a valid active scenario
        resp_wb = client.get(f"/exports/institutional-workbook.xlsx?project={code}")
        assert resp_wb.status_code == 200, (
            f"Institutional workbook export failed with active valid scenario: {resp_wb.text[:300]}")

        resp_csv = client.get(f"/exports/runtime-summary.csv?project={code}")
        assert resp_csv.status_code == 200, (
            f"Runtime CSV export failed with active valid scenario: {resp_csv.text[:300]}")

    def test_p_missing_active_scenario_cannot_export_base_silently(self):
        """When active_scenario_id references a non-existent scenario
        the export must return 400, not silently export Base economics."""
        from app.persistence.projects_repository import get_project_by_code

        client = _client()
        code = _create(client, "R5B Missing Scenario", "tuho")
        _run(client, code)

        rec = get_project_by_code(_USER_ID, code)
        # Force an invalid scenario_id into workspace state
        _force_active_scenario_id(rec.project_id, "nonexistent-scenario-uuid-r5b")

        resp = client.get(f"/exports/institutional-workbook.xlsx?project={code}")
        assert resp.status_code == 400, (
            f"Export silently returned {resp.status_code} for missing scenario — Gap B not closed")

        csv_resp = client.get(f"/exports/runtime-summary.csv?project={code}")
        assert csv_resp.status_code == 400, (
            f"CSV export silently returned {csv_resp.status_code} for missing scenario — Gap B not closed")

    def test_q_archived_active_scenario_cannot_export_base_silently(self):
        """When the active scenario is archived the export must return 400."""
        from app.persistence.projects_repository import get_project_by_code
        from app.persistence.scenarios_repository import add_scenario, select_scenario
        from app.persistence.workspace_repository import get_workspace_state
        from app.persistence.scenarios_repository import archive_scenario

        client = _client()
        code = _create(client, "R5B Archived Scenario", "tuho")
        _run(client, code)

        rec = get_project_by_code(_USER_ID, code)
        ws = get_workspace_state(_USER_ID, rec.project_id)
        sc = add_scenario(
            user_id=_USER_ID, project_id=rec.project_id,
            project_code=rec.project_code, scenario_name="R5B Archived Sc",
            parent_scenario_id=ws.active_scenario_id or "",
            base_input_set=dict(ws.draft_snapshot),
        )
        assert select_scenario(_USER_ID, rec.project_id, sc.scenario_id)
        # Archive the now-active scenario
        archive_scenario(_USER_ID, sc.scenario_id)

        resp = client.get(f"/exports/institutional-workbook.xlsx?project={code}")
        assert resp.status_code == 400, (
            f"Export silently returned {resp.status_code} for archived scenario — Gap B not closed")

        csv_resp = client.get(f"/exports/runtime-summary.csv?project={code}")
        assert csv_resp.status_code == 400, (
            f"CSV export silently returned {csv_resp.status_code} for archived scenario — Gap B not closed")

    def test_r_cross_project_active_scenario_cannot_export_base_silently(self):
        """When the active scenario belongs to a different project the export must
        return 400, not silently export Base economics for the wrong project."""
        from app.persistence.projects_repository import get_project_by_code
        from app.persistence.scenarios_repository import add_scenario
        from app.persistence.workspace_repository import get_workspace_state

        client = _client()
        # Project A
        code_a = _create(client, "R5B CrossProject A", "tuho")
        _run(client, code_a)
        # Project B — its scenario will be injected into Project A's workspace
        code_b = _create(client, "R5B CrossProject B", "tuho")
        _run(client, code_b)

        rec_b = get_project_by_code(_USER_ID, code_b)
        ws_b = get_workspace_state(_USER_ID, rec_b.project_id)
        sc_b = add_scenario(
            user_id=_USER_ID, project_id=rec_b.project_id,
            project_code=rec_b.project_code, scenario_name="R5B Cross Sc",
            parent_scenario_id=ws_b.active_scenario_id or "",
            base_input_set=dict(ws_b.draft_snapshot),
        )

        # Inject B's scenario into A's workspace
        rec_a = get_project_by_code(_USER_ID, code_a)
        _force_active_scenario_id(rec_a.project_id, sc_b.scenario_id)

        resp = client.get(f"/exports/institutional-workbook.xlsx?project={code_a}")
        assert resp.status_code == 400, (
            f"Export silently returned {resp.status_code} for cross-project scenario — Gap B not closed")

        csv_resp = client.get(f"/exports/runtime-summary.csv?project={code_a}")
        assert csv_resp.status_code == 400, (
            f"CSV export silently returned {csv_resp.status_code} for cross-project scenario — Gap B not closed")

    def test_s_base_scenario_base_no_stale_state_correction_b(self):
        """Valid active scenario → Base: selecting then deselecting a scenario
        must not leave stale state — the export returns to Base economics.
        Regression gate for Gap B fix."""
        from app.persistence.projects_repository import get_project_by_code
        from app.persistence.scenarios_repository import (
            add_scenario, list_scenarios, select_scenario,
        )
        from app.persistence.workspace_repository import get_workspace_state

        client = _client()
        code = _create(client, "R5B Round-trip Base", "tuho")
        _run(client, code)

        base_wb = _export_workbook(client, code)
        base_rs = _runtime_summary(base_wb)

        rec = get_project_by_code(_USER_ID, code)
        ws = get_workspace_state(_USER_ID, rec.project_id)
        sc = add_scenario(
            user_id=_USER_ID, project_id=rec.project_id,
            project_code=rec.project_code, scenario_name="R5B Round-trip Sc",
            parent_scenario_id=ws.active_scenario_id or "",
            base_input_set=dict(ws.draft_snapshot),
        )
        assert select_scenario(_USER_ID, rec.project_id, sc.scenario_id)

        # Scenario export must succeed (valid scenario)
        scen_resp = client.get(f"/exports/institutional-workbook.xlsx?project={code}")
        assert scen_resp.status_code == 200, (
            f"Scenario export failed unexpectedly: {scen_resp.text[:300]}")

        # Return to Base by clearing active scenario
        base_candidates = [
            s for s in list_scenarios(_USER_ID, rec.project_id)
            if getattr(s, "is_base", False)
        ]
        if base_candidates:
            assert select_scenario(_USER_ID, rec.project_id, base_candidates[0].scenario_id)
        else:
            _force_active_scenario_id(rec.project_id, None)

        restored_wb = _export_workbook(client, code)
        restored_rs = _runtime_summary(restored_wb)
        assert restored_rs["Project IRR"] == pytest.approx(base_rs["Project IRR"], rel=1e-9), (
            "Returning to Base produced different economics — stale scenario state regression")

    def test_t_both_xlsx_and_csv_share_fail_closed_contract(self):
        """Both institutional XLSX and runtime-summary CSV share the same
        fail-closed scenario authority contract (not just XLSX)."""
        from app.persistence.projects_repository import get_project_by_code

        client = _client()
        code = _create(client, "R5B CSV Fail Closed", "tuho")
        _run(client, code)

        rec = get_project_by_code(_USER_ID, code)
        _force_active_scenario_id(rec.project_id, "csv-fail-closed-missing-uuid")

        xlsx_resp = client.get(f"/exports/institutional-workbook.xlsx?project={code}")
        csv_resp = client.get(f"/exports/runtime-summary.csv?project={code}")
        assert xlsx_resp.status_code == 400, (
            f"XLSX: expected 400, got {xlsx_resp.status_code}")
        assert csv_resp.status_code == 400, (
            f"CSV: expected 400, got {csv_resp.status_code}")
